"""Генератор демо-архива: синтетические письма, акты, протоколы, сметы.

Нужен, чтобы проверять и показывать поиск, не трогая рабочие документы.
Реальные файлы компании в репозиторий не попадают никогда.
"""
from __future__ import annotations

import random
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

OUT = Path(__file__).resolve().parent.parent / "sample"

CONTRACTORS = ["ООО СтройМонтаж", "АО ЭнергоСервис", "ООО ПромГеоТех",
               "ООО ТеплоСети", "АО МостСтрой"]
OBJECTS = ["ПС-110/10-Северная", "КНС-4", "Котельная-3", "Эстакада-2"]
MATERIALS = ["арматура А500С", "щебень фр. 20-40", "бетон В25 W6",
             "кабель АВБбШв 4х95", "трубы ПЭ100 SDR17"]
WORKS = ["устройство монолитного ростверка", "прокладка кабельных линий",
         "монтаж металлоконструкций", "гидроизоляция фундаментов",
         "испытание трубопровода на прочность"]


def letter_text(num: int, date: str, contractor: str, obj: str) -> str:
    return (
        f"Исх. № {num}/ПТО от {date}\n\n"
        f"Генеральному директору\n{contractor}\n\n"
        f"Объект: {obj}\nШифр: {random.randint(100, 999)}-{random.randint(10, 99)}-ПЗ\n\n"
        f"Уважаемые коллеги!\n\n"
        f"Направляем в Ваш адрес уведомление о необходимости поставки "
        f"{random.choice(MATERIALS)} в объёме {random.randint(10, 400)} т "
        f"в срок до {date}. Задержка поставки влечёт смещение сроков по работам: "
        f"{random.choice(WORKS)}.\n\n"
        f"Просим подтвердить готовность и направить сопроводительную документацию.\n\n"
        f"Начальник ПТО\n"
    )


def act_text(num: int, date: str, contractor: str, obj: str) -> str:
    return (
        f"АКТ освидетельствования скрытых работ № {num}\n"
        f"г. Екатеринбург, {date}\n\n"
        f"Объект капитального строительства: {obj}\n"
        f"Подрядчик: {contractor}\n\n"
        f"Комиссия составила настоящий акт о том, что выполнены работы: "
        f"{random.choice(WORKS)}.\n"
        f"Применённые материалы: {random.choice(MATERIALS)}.\n"
        f"Работы выполнены в соответствии с проектной документацией и СП.\n"
        f"Разрешается производство последующих работ.\n"
    )


def protocol_text(num: int, date: str, contractor: str, obj: str) -> str:
    return (
        f"ПРОТОКОЛ № {num} совещания по объекту {obj}\n"
        f"от {date}\n\n"
        f"Присутствовали: представители заказчика, {contractor}.\n\n"
        f"ПОВЕСТКА:\n"
        f"1. Ход выполнения работ: {random.choice(WORKS)}.\n"
        f"2. Задержка поставки: {random.choice(MATERIALS)}.\n\n"
        f"РЕШИЛИ:\n"
        f"1. {contractor} обеспечить поставку в срок до {date}.\n"
        f"2. Начальнику ПТО подготовить откорректированный график.\n"
    )


def safe(name: str) -> str:
    """Убрать из имени объекта символы, недопустимые в имени файла."""
    for ch in ('/', chr(92), ':', '*', '?', '"', '<', '>', '|'):
        name = name.replace(ch, "-")
    return name


def write_docx(path: Path, text: str) -> None:
    import docx

    document = docx.Document()
    for line in text.split("\n"):
        document.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


# Встроенные шрифты PyMuPDF кириллицу не умеют, берём системный
CYRILLIC_FONT = Path("C:/Windows/Fonts/arial.ttf")


def write_pdf(path: Path, text: str) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    if CYRILLIC_FONT.exists():
        page.insert_font(fontname="cyr", fontfile=str(CYRILLIC_FONT))
        fontname = "cyr"
    else:
        fontname = "helv"
    page.insert_textbox(
        pymupdf.Rect(50, 50, 545, 780), text, fontsize=10, fontname=fontname
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()


def write_scan_pdf(path: Path) -> None:
    """PDF без текстового слоя — имитация скана подписанного письма."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.draw_rect(pymupdf.Rect(60, 60, 535, 760), color=(0.7, 0.7, 0.7))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    doc.close()


def write_xlsx(path: Path, obj: str, contractor: str) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Ведомость объёмов"
    ws.append(["№", "Наименование работ", "Ед.", "Кол-во", "Подрядчик", "Объект"])
    for i, work in enumerate(random.sample(WORKS, 4), 1):
        ws.append([i, work, "м3", random.randint(5, 500), contractor, obj])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def main() -> None:
    random.seed(42)
    if OUT.exists():
        print(f"Папка {OUT} уже существует — удалите её вручную, если нужен новый набор")
        return

    n = 0
    for year in (2024, 2025):
        for contractor in CONTRACTORS:
            obj = random.choice(OBJECTS)
            base = OUT / str(year) / contractor
            for i in range(1, 5):
                num = random.randint(100, 999)
                date = f"{random.randint(1, 28):02d}.{random.randint(1, 12):02d}.{year}"
                write_docx(base / "Письма" / f"Исх {num} от {date} {safe(obj)}.docx",
                           letter_text(num, date, contractor, obj))
                write_pdf(base / "Акты" / f"Акт АОСР {num} {date}.pdf",
                          act_text(num, date, contractor, obj))
                write_docx(base / "Протоколы" / f"Протокол {num} от {date}.docx",
                           protocol_text(num, date, contractor, obj))
                n += 3
            write_xlsx(base / "Ведомость объёмов работ.xlsx", obj, contractor)
            write_scan_pdf(base / "Письма" / f"Скан ответа {contractor}.pdf")
            (base / "Чертежи").mkdir(parents=True, exist_ok=True)
            (base / "Чертежи" / f"{safe(obj)} план сетей.dwg").write_bytes(b"AC1032 fake")
            n += 3

    print(f"Успех: создано {n} файлов в {OUT}")


if __name__ == "__main__":
    main()
