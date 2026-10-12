"""Сделать шаблон письма из любого вашего исходящего.

Берём отправленное письмо, выбрасываем его содержание и оставляем
вёрстку: логотип, шапку, таблицу «номер — адресат», подпись. Переменные
места помечаются подстановками вида {{number}}.

    python tools/make_template.py "Исх. РТП-1335.docx" templates/письмо.docx

Посмотреть, что размечено в готовом шаблоне:

    python tools/make_template.py --show templates/письмо.docx

Выровнять отступы в шапке, не пересобирая шаблон:

    python tools/make_template.py --align templates/письмо.docx

Старый .doc сначала пересохраните в .docx — Word умеет это сам.

Получившийся шаблон правится в Word как обычный документ: подстановки —
это просто текст. Главное не трогать фигурные скобки.
"""
from __future__ import annotations

import sys
from pathlib import Path

PLACEHOLDERS = {
    "number": "номер исходящего",
    "date": "дата",
    "recipient_position": "кому, должность",
    "recipient_org": "организация адресата",
    "recipient_person": "фамилия адресата",
    "subject": "тема",
    "greeting": "обращение",
    "intro": "первый абзац — основание, договор",
    "body": "текст письма",
    "signer_position": "должность подписанта",
    "signer_name": "фамилия подписанта",
}


def set_text(paragraph, text: str) -> None:
    """Заменить текст абзаца, сохранив оформление первого прогона."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for extra in runs[1:]:
        extra.text = ""


def drop(element) -> None:
    element.getparent().remove(element)


def strip_numbering(document) -> int:
    """Убрать из бланка нумерацию абзацев.

    В исходном письме был нумерованный список, и его разметка остаётся в
    пустых абзацах бланка: Word исправно рисует «1.» и «2.» на пустом
    месте. Нумерация бывает двух видов — прописанная прямо в абзаце и
    унаследованная от стиля, — и снимать нужно обе. Нумерация это
    свойство содержания, а не бланка: в письме она задаётся в форме.
    """
    from docx.oxml.ns import qn

    removed = 0
    for paragraph in document.paragraphs:
        pPr = paragraph._element.find(qn("w:pPr"))
        if pPr is not None:
            numPr = pPr.find(qn("w:numPr"))
            if numPr is not None:
                pPr.remove(numPr)
                removed += 1

        style = (paragraph.style.name or "").lower()
        if not paragraph.text.strip() and ("list" in style or "список" in style):
            paragraph.style = document.styles["Normal"]
            removed += 1
    return removed


def letterhead_paragraphs(document):
    """Абзацы шапки — всё до первой таблицы с номером и адресатом."""
    from docx.text.paragraph import Paragraph

    result = []
    for child in document.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "tbl":
            break
        if tag == "p":
            result.append(Paragraph(child, document))
    return result


EDGE_SPACE = " " + chr(0xa0) + chr(9)


def strip_edge_spaces(paragraph) -> bool:
    """Убрать пробелы по краям абзаца.

    В исходном бланке текст двигали пробелами: в одной строке их шестьдесят
    в начале, в другой восемь. При выравнивании по центру такие пробелы
    считаются частью строки и сдвигают её — реквизиты встают вразнобой.
    """
    runs = paragraph.runs
    if not runs:
        return False
    changed = False

    for run in runs:
        if not run.text:
            continue
        stripped = run.text.lstrip(EDGE_SPACE)
        if stripped != run.text:
            run.text = stripped
            changed = True
        if stripped:
            break

    for run in reversed(runs):
        if not run.text:
            continue
        stripped = run.text.rstrip(EDGE_SPACE)
        if stripped != run.text:
            run.text = stripped
            changed = True
        if stripped:
            break
    return changed


def align_letterhead(document) -> list[str]:
    """Выровнять отступы в шапке по одному значению.

    Строки шапки отодвинуты вправо, чтобы не налезать на логотип, но
    отступ у них разный — у одной строки 5,0 см, у остальных 5,5. При
    выравнивании по центру каждая строка центруется в своём поле, и
    реквизиты визуально разъезжаются.
    """
    from collections import Counter

    paragraphs = [p for p in letterhead_paragraphs(document) if p.text.strip()]
    if not paragraphs:
        return []

    report = []
    spaced = sum(1 for p in paragraphs if strip_edge_spaces(p))
    if spaced:
        report.append(f"убраны пробелы по краям строк шапки: {spaced}")

    indents = [p.paragraph_format.left_indent for p in paragraphs]
    known = [int(value) for value in indents if value is not None]
    if not known:
        return report

    common = Counter(known).most_common(1)[0][0]
    from docx.shared import Emu

    changed = 0
    for paragraph in paragraphs:
        current = paragraph.paragraph_format.left_indent
        if current is None or int(current) != common:
            paragraph.paragraph_format.left_indent = Emu(common)
            paragraph.paragraph_format.right_indent = None
            changed += 1

    if changed:
        report.append(f"выровнены отступы в шапке: {changed} строк(и)"
                      f" по {Emu(common).cm:.2f} см")
    return report


def build(source: str, target: str) -> list[str]:
    import docx

    document = docx.Document(source)
    report: list[str] = []

    if document.tables:
        head = document.tables[0]
        cells = head.rows[0].cells
        set_text(cells[0].paragraphs[0], "Исх. {{number}} от {{date}}")
        report.append("таблица шапки: номер слева, адресат справа")

        if len(cells) > 1:
            wanted = ["{{recipient_position}}", "{{recipient_org}}",
                      "{{recipient_person}}"]
            paragraphs = cells[1].paragraphs
            for i, placeholder in enumerate(wanted):
                if i < len(paragraphs):
                    set_text(paragraphs[i], placeholder)
                else:
                    cells[1].add_paragraph(placeholder)
            for extra in paragraphs[len(wanted):]:
                drop(extra._element)

    # Остальные таблицы — приложения конкретного письма, в бланке лишние
    for table in document.tables[1:]:
        drop(table._element)
        report.append("удалена таблица приложения")

    state = "before"
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        low = text.lower()

        if low.startswith("тема"):
            set_text(paragraph, "Тема: «{{subject}}»")
            report.append("тема")
        elif low.startswith("уважаем"):
            set_text(paragraph, "{{greeting}}")
            state = "body"
            report.append("обращение")
        elif low.startswith("с уважением"):
            state = "signature"
        elif state == "signature" and text:
            set_text(paragraph, "{{signer_position}}\t\t{{signer_name}}")
            state = "done"
            report.append("подпись")
        elif state == "body" and text:
            # первый абзац письма — всегда одно и то же основание
            set_text(paragraph, "{{intro}}")
            state = "intro_done"
            report.append("преамбула (первый абзац)")
        elif state == "intro_done" and text:
            set_text(paragraph, "{{body}}")
            state = "body_done"
            report.append("текст письма")
        elif state == "body_done" and text:
            drop(paragraph._element)

    if state == "body_done":
        report.append("подпись не найдена — допишите {{signer_position}} вручную")

    report += align_letterhead(document)
    numbered = strip_numbering(document)
    if numbered:
        report.append(f"снята нумерация с абзацев: {numbered}")

    Path(target).parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return report


def show(template: str) -> int:
    """Показать, что в шаблоне размечено — когда письмо выходит не таким."""
    import docx

    document = docx.Document(template)
    print(f"Шаблон: {template}")
    print()
    for i, paragraph in enumerate(document.paragraphs):
        text = paragraph.text.strip()
        if text:
            print(f"  {i:3} {text[:90]}")
    for ti, table in enumerate(document.tables):
        for row in table.rows:
            for cell in row.cells:
                joined = " | ".join(p.text.strip() for p in cell.paragraphs
                                    if p.text.strip())
                if joined:
                    print(f"  [таблица {ti}] {joined[:90]}")
    print()
    found = {key for key in PLACEHOLDERS
             if "{{" + key + "}}" in docx.Document(template).element.xml}
    missing = [key for key in PLACEHOLDERS if key not in found]
    print("Найдены подстановки:", ", ".join(sorted(found)) or "ни одной")
    if missing:
        print("Отсутствуют:", ", ".join(missing))
    return 0


def align_only(template: str) -> int:
    """Поправить отступы в готовом шаблоне, не пересобирая его."""
    import docx

    document = docx.Document(template)
    report = align_letterhead(document)
    strip = strip_numbering(document)
    if strip:
        report.append(f"снята нумерация с абзацев: {strip}")
    if not report:
        print("Менять нечего: отступы в шапке уже одинаковые")
        return 0
    document.save(template)
    print(f"Сделал: {template}")
    for line in report:
        print(f"  {line}")
    return 0


def main() -> int:
    if len(sys.argv) == 3 and sys.argv[1] == "--show":
        return show(sys.argv[2])
    if len(sys.argv) == 3 and sys.argv[1] == "--align":
        return align_only(sys.argv[2])
    if len(sys.argv) != 3:
        print(__doc__)
        return 1
    source, target = sys.argv[1], sys.argv[2]
    if not Path(source).exists():
        print(f"Неуспех: нет файла {source}")
        return 1
    if Path(source).suffix.lower() != ".docx":
        print("Неуспех: нужен .docx. Старый .doc пересохраните в Word")
        return 1

    report = build(source, target)
    print(f"Сделал: шаблон {target}")
    for line in report:
        print(f"  размечено: {line}")
    print()
    print("Пропишите его в конфиг:")
    print(f"  letterhead:\n    template: \"{target}\"")
    print()
    print("Доступные подстановки:")
    for key, meaning in PLACEHOLDERS.items():
        print(f"  {{{{{key}}}}}".ljust(28) + meaning)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
