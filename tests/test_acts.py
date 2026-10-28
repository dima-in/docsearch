"""Акты комиссии: шаблон из бланка, комиссия, номер, пакет.

Люди и организации здесь выдуманы: настоящие акты в репозиторий не идут.
"""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from fastapi.testclient import TestClient

from docsearch import acts, db, indexer
from docsearch.config import Config, Root
from docsearch.web import create_app

NL = chr(10)


def blank(path: Path) -> Path:
    """Бланк, устроенный как настоящий: всё в ячейке таблицы, прочерки,
    выравнивание пробелами, три пары «строка — подпись под ней»."""
    document = Document()
    table = document.add_table(rows=3, cols=1)
    head, number, main = (row.cells[0] for row in table.rows)

    p = head.paragraphs[0]
    p.add_run("Строительная организация ")
    p.add_run("__________")
    bold = p.add_run("_____")
    bold.bold = bold.underline = True
    head.add_paragraph("Строительство «Жилой дом на Лесной улице»")
    head.add_paragraph("(наименование и месторасположение)")

    p = number.paragraphs[0]
    for text in ("АКТ ", "№ ", "_____"):
        p.add_run(text).bold = True

    main.paragraphs[0].text = "г. Москва" + " " * 80 + "« ___ » _____ 2026 г."
    main.add_paragraph("")
    main.add_paragraph("    Комиссия в составе:")
    for _ in range(3):
        line = main.add_paragraph().add_run(" " * 40 + "____________")
        line.bold = line.italic = line.underline = True
        main.add_paragraph("(должности, фамилии, инициалы)")
    main.add_paragraph("")
    main.add_paragraph("составили настоящий Акт в том, что ООО ____ передает ____")
    main.add_paragraph("__________________________________")
    main.add_paragraph("__________________________________.")
    main.add_paragraph("")
    main.add_paragraph("(Основание Договор № __________)")
    main.add_paragraph("")
    sign = main.add_paragraph().add_run("_________" + " " * 10 + "________")
    sign.bold = True
    main.add_paragraph(" " * 30 + "(подпись)")
    out = path / "бланк.docx"
    document.save(out)
    return out


@pytest.fixture
def settings(tmp_path: Path) -> acts.ActSettings:
    template = tmp_path / "акт.docx"
    acts.make_template(blank(tmp_path)).save(template)
    return acts.ActSettings(template=str(template), organization="ООО «Альфа»")


def text_of(blob: bytes) -> str:
    """Весь текст документа, включая ячейки: бланк свёрстан таблицей."""
    document = Document(io.BytesIO(blob))
    return NL.join(Paragraph(p, None).text
                   for p in document.element.body.iter(qn("w:p")))


MEMBERS = [
    {"position": "Руководитель проекта", "org": "ООО «Альфа»", "name": "Петров И.И."},
    {"position": "Начальник участка", "org": "ООО «Бета»", "name": "Сидоров П.П."},
    {"position": "Представитель", "org": "ООО «Гамма»", "absent": True},
]


# — разбор комиссии ———————————————————————————————————————————————

def test_member_line_with_name():
    m = acts.parse_member("Руководитель проекта ООО «Альфа»  Петров И. И.___ ____ ______")
    assert (m.position, m.org, m.name, m.absent) == \
        ("Руководитель проекта", "ООО «Альфа»", "Петров И.И.", False)


def test_absent_member():
    m = acts.parse_member("Представитель ООО «Гамма»_ отсутствовал, акт составлен "
                          "в одностороннем порядке ______")
    assert m.org == "ООО «Гамма»" and m.absent and not m.name


def test_member_without_name_and_individual():
    assert acts.parse_member("Представитель ООО «Дельта» ________").name == ""
    ip = acts.parse_member("ИП Кузнецов А.В. ________")
    assert ip.org == "ИП Кузнецов А.В." and ip.position == ""
    quoted = acts.parse_member("Подрядчик ИП «Орлов Р.В.»")
    assert quoted.org == "ИП «Орлов Р.В.»" and quoted.position == "Подрядчик"


def test_captions_and_dashes_are_not_members():
    assert acts.parse_member("(должности, фамилии, инициалы)") is None
    assert acts.parse_member("______________________") is None


def test_commission_stops_before_the_act_text():
    """В тексте акта организации тоже названы — это не комиссия."""
    text = NL.join(["АКТ № 5", "Комиссия в составе:",
                    "Руководитель проекта ООО «Альфа» Петров И.И.____",
                    "(должность, фамилия, инициалы)",
                    "составили настоящий Акт в том, что ООО «Бета» передало"])
    members = acts.commission_of(text)
    assert [m.org for m in members] == ["ООО «Альфа»"]


# — папка актов ——————————————————————————————————————————————————

ACT = NL.join(["АКТ № {n}", "г. Москва « 1 » марта 2026 г.", "Комиссия в составе:",
               "Руководитель проекта ООО «Альфа» {lead}___",
               "(должность, фамилия, инициалы)",
               "Начальник участка ООО «{other}» Сидоров П.П.___",
               "составили настоящий Акт в том, что работы выполнены"])


@pytest.fixture
def archive(tmp_path: Path):
    root = tmp_path / "arc"
    folder = root / "АКТЫ"
    (folder / "Для дверей").mkdir(parents=True)
    (folder / "Акт № 7.txt").write_text(
        ACT.format(n=7, lead="Петров И. И.", other="Бета"), encoding="utf-8")
    (folder / "Акт № 8 передача.txt").write_text(
        ACT.format(n=8, lead="Петров И.И.", other="БЕТА"), encoding="utf-8")
    (folder / "Для дверей" / "Акт № 1.txt").write_text(
        ACT.format(n=1, lead="Петров И.И.", other="Бета"), encoding="utf-8")
    # другая серия с большими номерами — в счёт идти не должна
    (root / "Акты ЗОС").mkdir()
    (root / "Акты ЗОС" / "Акт № 300.txt").write_text(
        ACT.format(n=300, lead="Петров И.И.", other="Бета"), encoding="utf-8")
    cfg = Config(roots=[Root(label="ПТО", path=str(root))],
                 db=str(tmp_path / "index.db"), own_org="ООО «Альфа»")
    conn = db.connect(cfg.db)
    indexer.run(conn, cfg)
    yield conn, cfg
    conn.close()


def test_last_act_is_the_highest_number_not_the_latest_date(archive, tmp_path):
    """В тексте акта № 3 упомянута дата позже, чем в № 8, — последним
    всё равно остаётся № 8: нумерация сквозная."""
    conn, cfg = archive
    late = ACT.format(n=3, lead="Петров И.И.", other="Бета").replace(
        "« 1 » марта 2026 г.", "« 1 » декабря 2026 г.")
    (Path(cfg.roots[0].path) / "АКТЫ" / "Акт № 3.txt").write_text(late, encoding="utf-8")
    indexer.run(conn, cfg)
    _, last = acts.directory(conn, "АКТЫ")
    assert last["number"] == 8


def test_next_number_counts_only_the_acts_folder(archive):
    conn, _ = archive
    assert acts.next_number(conn, "АКТЫ") == 9


def test_directory_merges_spellings_and_puts_ours_first(archive):
    conn, _ = archive
    people, last = acts.directory(conn, "АКТЫ", "ООО «Альфа»")
    assert people[0]["name"] == "Петров И.И." and people[0]["own"]
    beta = [p for p in people if acts.letters.org_key(p["org"]) == "бета"]
    assert len(beta) == 1                  # «Бета» и «БЕТА» — одна организация
    assert last["number"] == 8
    assert [m["name"] for m in last["members"]] == ["Петров И.И.", "Сидоров П.П."]


# — шаблон и заполнение ————————————————————————————————————————————

def test_template_marks_everything(settings):
    found = acts.placeholders(Document(settings.template))
    for token in ("organization", "number", "day", "month", "year", "#commission",
                  "line", "body", "?appendix", "?basis", "#signatures", "sign"):
        assert token in found, token


def test_commission_goes_both_to_the_top_and_to_signatures(settings):
    name, blob = acts.build({"number": 12, "date": "2026-03-05", "members": MEMBERS,
                             "body": "работы выполнены"}, settings)[0]
    text = text_of(blob)
    assert name == "Акт № 12.docx"
    assert text.count("Петров И.И.") == 2
    assert "Представитель ООО «Гамма» отсутствовал, акт составлен" in text
    assert "« 05 » марта 2026 г." in text
    assert "АКТ № 12" in text
    assert "{{" not in text
    # прочерки бланка ушли, остался только текст
    assert "передает ____" not in text


def test_empty_optional_fields_leave_no_lines(settings):
    blob = acts.build({"number": 1, "members": MEMBERS[:1], "body": "x"}, settings)[0][1]
    text = text_of(blob)
    assert "Приложение" not in text
    assert "Основание" not in text


def test_filled_optional_fields_and_multiline_body(settings):
    blob = acts.build({"number": 1, "members": MEMBERS[:1],
                       "body": "первая строка" + NL + NL + "вторая строка",
                       "appendix": "фото на 2 листах", "basis": "Договор № 5"},
                      settings)[0][1]
    lines = text_of(blob).split(NL)
    assert "составили настоящий Акт в том, что первая строка" in lines
    assert "вторая строка" in lines
    assert "Приложение: фото на 2 листах" in lines
    assert "(Основание: Договор № 5)" in lines


def test_organization_keeps_its_formatting(settings):
    """Название в бланке жирное и подчёркнутое — так и должно остаться."""
    blob = acts.build({"number": 1, "members": MEMBERS[:1], "body": "x"}, settings)[0][1]
    document = Document(io.BytesIO(blob))
    run = next(r for p in document.tables[0].rows[0].cells[0].paragraphs
               for r in p.runs if "Альфа" in r.text)
    assert run.bold and run.underline


def test_signature_line_is_a_tab_not_spaces(settings):
    """Линия подписи — табуляция с подчёркиванием: длинная фамилия не
    сдвигает её, как сдвигали пробелы в бланке."""
    blob = acts.build({"number": 1, "members": MEMBERS[:1], "body": "x"}, settings)[0][1]
    document = Document(io.BytesIO(blob))
    sign = next(p for p in document.tables[0].rows[2].cells[0].paragraphs
                if p.text.startswith("ООО «Альфа»\t"))
    leaders = [t.leader for t in sign.paragraph_format.tab_stops]
    assert any(str(leader).startswith("LINES") for leader in leaders)


# — пакет ———————————————————————————————————————————————————————

def test_batch_one_act_per_row_numbers_in_a_row(settings):
    files = acts.build({
        "number": 52, "members": MEMBERS,
        "body": "произведен осмотр работ по КС-2 № {КС-2} шифр {шифр}. {замечания}",
        "rows": "1/6\tКК3\tНет заземления лотков" + NL + "1/19;ВК1;",
    }, settings)
    assert [n for n, _ in files] == ["Акт № 52.docx", "Акт № 53.docx"]
    second = text_of(files[1][1])
    assert "КС-2 № 1/19 шифр ВК1." in second
    assert "{" not in second


def test_batch_without_rows_says_what_is_missing(settings):
    with pytest.raises(ValueError, match="шифр"):
        acts.build({"number": 1, "members": MEMBERS, "body": "шифр {шифр}"}, settings)


def test_fields_are_found_in_order_and_once():
    assert acts.fields({"body": "{a} и {b}, снова {a}", "subject": "{c}"}) == ["a", "b", "c"]
    assert acts.fields({"body": "{{number}} — это не поле"}) == []


# — через сервер ———————————————————————————————————————————————————

def test_api_draft_and_build(archive, settings):
    conn, cfg = archive
    cfg.acts = {"template": settings.template, "folder": "АКТЫ"}
    client = TestClient(create_app(cfg))

    draft = client.get("/api/act/draft").json()
    assert draft["ready"] and draft["number"] == 9 and draft["number_after"] == 8
    assert draft["last"]["number"] == 8

    one = client.post("/api/act", json={"number": 9, "members": MEMBERS,
                                        "body": "x", "subject": "передача"})
    assert one.status_code == 200
    assert "docx" in one.headers["content-disposition"]

    pack = client.post("/api/act", json={"number": 9, "members": MEMBERS,
                                         "body": "{шифр}", "rows": "А" + NL + "Б"})
    assert pack.headers["content-type"] == "application/zip"
    names = zipfile.ZipFile(io.BytesIO(pack.content)).namelist()
    assert names == ["Акт № 9.docx", "Акт № 10.docx"]


def test_api_without_template_explains(archive):
    conn, cfg = archive
    client = TestClient(create_app(cfg))
    assert client.get("/api/act/draft").json()["ready"] is False
    response = client.post("/api/act", json={"number": 1, "body": "x"})
    assert response.status_code == 400
    assert "act-template" in response.text
