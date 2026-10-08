"""Старые форматы Office и порядок чтения документа."""
from __future__ import annotations

from pathlib import Path

import pytest

from docsearch import letters, office
from docsearch.extract.word import extract_docx


def test_document_is_read_in_order(tmp_path: Path):
    """Таблица сверху должна попасть в начало текста, а не в конец.

    В бланке письма номер и адресат лежат в таблице. При раздельном
    чтении абзацев и таблиц они оказывались за пределами шапки, и по ним
    не разбирались ни номер, ни получатель.
    """
    import docx

    path = tmp_path / "письмо.docx"
    document = docx.Document()
    document.add_paragraph("ООО «ФБ-СТРОЙ»")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Исх. РТП-1335 от 06.10.2026 г."
    table.rows[0].cells[1].text = "Генеральному директору"
    document.add_paragraph("Тема: «Ответ»")
    document.save(str(path))

    text = extract_docx(path).text
    assert text.index("РТП-1335") < text.index("Тема:")


def test_addressee_found_inside_a_table_cell():
    """Ячейки строки склеиваются через « | » — для разбора это перенос."""
    text = ("Исх. РТП-1335 от 06.10.2026 г. | Генеральному директору "
            + chr(10) + "ООО «Мосренстрой-6»" + chr(10) + "Севрюкову Е.В.")
    found = letters.parse_addressee(text)
    assert found["position"] == "Генеральному директору"
    assert found["org"] == "ООО «Мосренстрой-6»"
    assert found["person"] == "Севрюкову Е.В."


def test_kind_for_known_formats():
    assert office.kind_for(".doc") == "word"
    assert office.kind_for(".xls") == "excel"
    assert office.kind_for(".pdf") is None


def test_converted_extension_map():
    assert office.CONVERTED[".doc"] == ".docx"
    assert office.CONVERTED[".xls"] == ".xlsx"


def test_empty_batch_does_nothing():
    assert office.convert_batch([], "word") == {}


def test_check_explains_what_is_missing(monkeypatch):
    monkeypatch.setattr(office, "available", lambda kind="word": False)
    with pytest.raises(office.OfficeUnavailable) as exc:
        office.check({"word"})
    assert "Word" in str(exc.value)
    assert "Office" in str(exc.value)
