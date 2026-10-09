"""Генератор исходящих писем."""
from __future__ import annotations

import zipfile
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from docsearch import db, indexer, letters
from docsearch.config import Config, Root
from docsearch.web import create_app

HEAD = {
    "name": "ООО «ФБ-СТРОЙ»",
    "legal_address": "141014, г. Мытищи, ул. Веры Волошиной, д. 14",
    "inn": "5029172308", "kpp": "502901001",
    "phone": "8(499)649-00-50", "email": "info@fbstroy.ru",
    "signer_position": "Начальник ПТО", "signer_name": "Иноземцев Д. А.",
    "number_prefix": "РТП",
}


@pytest.fixture
def env(tmp_path: Path):
    root = tmp_path / "arc" / "Переписка"
    root.mkdir(parents=True)
    (root / "Исх РТП-175.txt").write_text(
        "Исх. РТП-175 от 19.03.2024 г.\nООО «Маренго»", encoding="utf-8")
    (root / "Исх РТП-9.txt").write_text(
        "Исх. РТП-9 от 01.02.2024 г.\nООО «Высота»", encoding="utf-8")
    cfg = Config(roots=[Root(label="ПТО", path=str(tmp_path / "arc"))],
                 db=str(tmp_path / "index.db"), letterhead=HEAD)
    conn = db.connect(cfg.db)
    indexer.run(conn, cfg)
    yield conn, cfg
    conn.close()


def test_parse_number():
    assert letters.parse_number("РТП-175") == ("РТП", 175)
    assert letters.parse_number("270/ПТО") == ("", 270)
    assert letters.parse_number("без цифр") is None
    assert letters.parse_number("") is None


def test_next_number_continues_the_archive(env):
    """Журнал исходящих никто не ведёт — максимум берём из самих писем."""
    conn, _ = env
    assert letters.next_number(conn, "РТП") == "РТП-176"


def test_next_number_ignores_other_prefixes(env):
    conn, _ = env
    assert letters.next_number(conn, "ВХ") == "ВХ-1"


def test_known_recipients_come_from_archive(env):
    conn, _ = env
    assert "ООО «Маренго»" in letters.known_recipients(conn)


def test_ru_date():
    assert letters.ru_date("2024-03-19") == "19 марта 2024 г."
    assert letters.ru_date("мусор").endswith("г.")


def test_letterhead_skips_empty_fields():
    head = letters.Letterhead.from_config({"name": "ООО «Тест»"})
    assert head.header_lines() == ["ООО «Тест»"]


def test_docx_contains_the_fields():
    head = letters.Letterhead.from_config(HEAD)
    blob = letters.build_docx({
        "number": "РТП-176", "date": "2026-10-06",
        "recipient_org": "ООО «Маренго»", "recipient_person": "Белякову А. В.",
        "subject": "Нарушение сроков", "reference": "РТП-145 от 18.01.2024",
        "body": "Просим устранить замечания.",
    }, head)
    with zipfile.ZipFile(BytesIO(blob)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    for expected in ("РТП-176", "ООО «Маренго»", "Белякову А. В.",
                     "Нарушение сроков", "Просим устранить замечания.",
                     "Начальник ПТО", "6 октября 2026 г.", "РТП-145"):
        assert expected in xml, expected


def test_file_name_follows_the_archive_convention():
    """В папке исходящих принято «Исх. РТП-1337 ПКС Инжиниринг (о чём)»."""
    assert letters.file_name({
        "number": "РТП-1338", "recipient_org": "ООО «ПКС-Инжиниринг»",
        "subject": "ответ УКМ-299", "date": "2026-10-09",
    }) == "Исх. РТП-1338 ПКС-Инжиниринг (ответ УКМ-299).docx"


def test_file_name_without_subject():
    assert letters.file_name({
        "number": "РТП-1340", "recipient_org": "ООО «М-СТРОЙ»",
    }) == "Исх. РТП-1340 М-СТРОЙ.docx"


def test_file_name_without_organization():
    assert letters.file_name({
        "number": "РТП-1341", "subject": "общее совещание",
    }) == "Исх. РТП-1341 (общее совещание).docx"


def test_file_name_has_no_forbidden_characters():
    name = letters.file_name({"subject": "Смета 1/2 по объекту: раздел ОВ"})
    for bad in ("/", chr(92), ":", "*", "?", '"', "<", ">", "|"):
        assert bad not in name


def test_file_name_falls_back():
    assert letters.file_name({}) == "Письмо.docx"


def test_short_org_strips_the_legal_form():
    assert letters.short_org("ООО «ПКС-Инжиниринг»") == "ПКС-Инжиниринг"
    assert letters.short_org("ГУП «Мосводосток»") == "Мосводосток"


def test_draft_endpoint(env):
    conn, cfg = env
    client = TestClient(create_app(cfg))
    draft = client.get("/api/letter/draft").json()
    assert draft["number"] == "РТП-176"
    assert draft["signer_name"] == "Иноземцев Д. А."
    orgs = [item["org"] for item in draft["recipients"]]
    assert "ООО «Маренго»" in orgs


def test_build_endpoint_returns_docx(env):
    conn, cfg = env
    client = TestClient(create_app(cfg))
    response = client.post("/api/letter", json={
        "number": "РТП-176", "date": "2026-10-06",
        "recipient_org": "ООО «Маренго»", "subject": "Проверка",
        "body": "Текст письма.",
    })
    assert response.status_code == 200
    assert response.headers["content-type"].endswith("wordprocessingml.document")
    assert zipfile.ZipFile(BytesIO(response.content)).namelist()


def test_build_refuses_without_letterhead(tmp_path: Path):
    """Письмо без бланка отправлять нельзя, и молчать об этом тоже."""
    cfg = Config(roots=[], db=str(tmp_path / "index.db"))
    conn = db.connect(cfg.db)
    conn.close()
    client = TestClient(create_app(cfg))
    response = client.post("/api/letter", json={"recipient_org": "ООО «Х»"})
    assert response.status_code == 400
    assert "letterhead" in response.json()["detail"]


def test_register_writes_exact_attributes(env):
    """Атрибуты сгенерированного письма не угадываются, а берутся как есть."""
    conn, cfg = env
    path = str(Path(cfg.roots[0].path) / "Переписка" / "Исх РТП-176.docx")
    letters.register(conn, path, {
        "number": "РТП-176", "date": "2026-10-06",
        "recipient_org": "ООО «Маренго»", "subject": "Проверка",
    })
    saved = db.get_override(conn, path)
    assert saved["doc_number"] == "РТП-176"
    assert saved["counterparty"] == "ООО «Маренго»"
    assert saved["section"] == "переписка"


def make_template(path: Path) -> Path:
    """Бланк для теста: своего в репозитории нет и быть не должно."""
    import docx

    document = docx.Document()
    document.add_paragraph("ООО «Образец»")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].text = "Исх. {{number}} от {{date}}"
    right = table.rows[0].cells[1]
    right.paragraphs[0].text = "{{recipient_position}}"
    right.add_paragraph("{{recipient_org}}")
    right.add_paragraph("{{recipient_person}}")
    document.add_paragraph("Тема: «{{subject}}»")
    document.add_paragraph("{{greeting}}")
    document.add_paragraph("{{body}}")
    document.add_paragraph("{{signer_position}}\t\t{{signer_name}}")
    document.save(str(path))
    return path


def read_all(blob: bytes) -> list[str]:
    import docx

    document = docx.Document(BytesIO(blob))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                lines += [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    return lines


def test_template_is_filled(tmp_path: Path):
    template = make_template(tmp_path / "бланк.docx")
    head = letters.Letterhead.from_config({**HEAD, "template": str(template)})
    blob = letters.render({
        "number": "РТП-1336", "date": "2026-10-07",
        "recipient_position": "Генеральному директору",
        "recipient_org": "ООО «Мосренстрой-6»",
        "recipient_person": "Севрюкову Е. В.",
        "subject": "Ответ на УКМ-300", "greeting": "Уважаемый Евгений Владимирович!",
        "body": "Первый абзац.",
    }, head)
    lines = read_all(blob)
    assert "Исх. РТП-1336 от 7 октября 2026 г." in lines
    assert "ООО «Мосренстрой-6»" in lines
    assert "Тема: «Ответ на УКМ-300»" in lines
    assert "Первый абзац." in lines
    assert not any("{{" in line for line in lines)


def test_multiline_body_becomes_several_paragraphs(tmp_path: Path):
    template = make_template(tmp_path / "бланк.docx")
    head = letters.Letterhead.from_config({**HEAD, "template": str(template)})
    blob = letters.render({"recipient_org": "ООО «Х»",
                           "body": "Первый.\nВторой.\nТретий."}, head)
    lines = read_all(blob)
    for expected in ("Первый.", "Второй.", "Третий."):
        assert expected in lines


def test_signer_falls_back_to_letterhead(tmp_path: Path):
    template = make_template(tmp_path / "бланк.docx")
    head = letters.Letterhead.from_config({**HEAD, "template": str(template)})
    lines = read_all(letters.render({"recipient_org": "ООО «Х»"}, head))
    assert any("Иноземцев Д. А." in line for line in lines)


def test_missing_fields_leave_no_placeholders(tmp_path: Path):
    """Незаполненное поле должно исчезнуть, а не остаться скобками в письме."""
    template = make_template(tmp_path / "бланк.docx")
    head = letters.Letterhead.from_config({**HEAD, "template": str(template)})
    lines = read_all(letters.render({"recipient_org": "ООО «Х»"}, head))
    assert not any("{{" in line or "}}" in line for line in lines)


def test_falls_back_when_template_missing(tmp_path: Path):
    head = letters.Letterhead.from_config(
        {**HEAD, "template": str(tmp_path / "нет.docx")})
    lines = read_all(letters.render({"recipient_org": "ООО «Х»",
                                     "body": "Текст."}, head))
    assert "ООО «ФБ-СТРОЙ»" in lines     # программная вёрстка сработала


def test_template_builder_marks_the_places(tmp_path: Path):
    """Инструмент из tools/ размечает отправленное письмо в шаблон."""
    import importlib.util

    source = tmp_path / "письмо.docx"
    import docx

    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].text = "Исх. РТП-1335 от 06.10.2026 г."
    right = table.rows[0].cells[1]
    right.paragraphs[0].text = "Генеральному директору"
    right.add_paragraph("ООО «Мосренстрой-6»")
    right.add_paragraph("Севрюкову Е.В.")
    document.add_paragraph("Тема: «Ответ на письмо»")
    document.add_paragraph("Уважаемый Евгений Владимирович!")
    document.add_paragraph("Между сторонами заключен договор подряда.")
    document.add_paragraph("Второй абзац письма.")
    document.add_paragraph("С уважением,")
    document.add_paragraph("Руководитель проекта\t\tСтерхов Д.А.")
    document.save(str(source))

    spec = importlib.util.spec_from_file_location(
        "make_template",
        Path(__file__).resolve().parent.parent / "tools" / "make_template.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    target = tmp_path / "шаблон.docx"
    module.build(str(source), str(target))

    result = docx.Document(str(target))
    lines = [p.text.strip() for p in result.paragraphs if p.text.strip()]
    assert "Тема: «{{subject}}»" in lines
    assert "{{greeting}}" in lines
    assert "{{body}}" in lines
    assert "Второй абзац письма." not in lines     # содержание выброшено
    assert any("{{signer_name}}" in line for line in lines)


def test_addressee_parsed_from_letter_body():
    """Блок «кому» в исходящем — это готовая карточка адресата."""
    text = (chr(10).join([
        "Исх. РТП-1335 от 06.10.2026 г.",
        "Генеральному директору",
        "ООО «Мосренстрой-6»",
        "Севрюкову Е.В.",
        "Тема: «Ответ»",
    ]))
    found = letters.parse_addressee(text)
    assert found["position"] == "Генеральному директору"
    assert found["org"] == "ООО «Мосренстрой-6»"
    assert found["person"] == "Севрюкову Е.В."


def test_addressee_absent():
    assert letters.parse_addressee("Просто текст без адресата") is None
    assert letters.parse_addressee("") is None


def test_contacts_are_built_from_correspondence(env):
    conn, cfg = env
    conn.execute(
        "UPDATE documents SET section = 'переписка', doc_type = 'письмо'")
    conn.commit()
    conn.execute(
        "DELETE FROM doc_fts WHERE rowid = (SELECT MIN(id) FROM documents)")
    doc_id = conn.execute("SELECT MIN(id) id FROM documents").fetchone()["id"]
    body = chr(10).join(["Генеральному директору", "ООО «Маренго»",
                         "Белякову А. В."])
    conn.execute("INSERT INTO doc_fts (rowid, name, body, lemmas)"
                 " VALUES (?,?,?,?)", (doc_id, "письмо", body, body))
    conn.commit()

    rows = letters.rebuild_contacts(conn)
    found = {r["org"]: r for r in rows}
    assert found["ООО «Маренго»"]["person"] == "Беляков А. В." or            found["ООО «Маренго»"]["person"] == "Белякову А. В."
    assert found["ООО «Маренго»"]["position"] == "Генеральному директору"


def test_contacts_endpoint(env):
    conn, cfg = env
    client = TestClient(create_app(cfg))
    assert client.post("/api/letter/contacts").status_code == 200


def test_letterhead_includes_ogrn_and_fax():
    """Реквизиты идут одной строкой, как в настоящем бланке."""
    head = letters.Letterhead.from_config({
        "name": "ООО «ФБ-СТРОЙ»", "ogrn": "1135029002547",
        "inn": "5029172308", "kpp": "772801001",
        "phone": "8(499)649-00-50", "fax": "8(495)785-39-30",
        "email": "info@fbstroy.ru",
    })
    lines = head.header_lines()
    assert "ОГРН 1135029002547 ИНН 5029172308 КПП 772801001" in lines
    assert any("Факс: 8(495)785-39-30" in line for line in lines)


def test_one_bad_number_does_not_poison_numbering():
    """Скан с лишней цифрой задирал счётчик на порядок: РТП-19095."""
    assert letters.highest_sane([1300, 1310, 1320, 1330, 1335, 19095]) == 1335


def test_real_jump_in_numbering_is_kept():
    """Если нумерация действительно ушла вверх, там плотная группа."""
    assert letters.highest_sane([100, 19000, 19050, 19080, 19095]) == 19095


def test_small_sample_is_trusted_as_is():
    assert letters.highest_sane([9, 175]) == 175
    assert letters.highest_sane([5]) == 5
    assert letters.highest_sane([]) == 0


def test_template_path_is_relative_to_config(tmp_path: Path):
    """Бланк должен находиться независимо от того, откуда запущено."""
    from docsearch import config as config_mod

    (tmp_path / "templates").mkdir()
    make_template(tmp_path / "templates" / "письмо.docx")
    cfg_file = tmp_path / "config.local.yaml"
    cfg_file.write_text(
        "roots:\n  - path: './arc'\nindex:\n  db: 'index.db'\n"
        "letterhead:\n  name: 'ООО «Тест»'\n"
        "  template: 'templates/письмо.docx'\n",
        encoding="utf-8")
    cfg = config_mod.load(cfg_file)
    assert Path(cfg.letterhead["template"]).exists()


def test_contract_is_mined_from_the_letter():
    """Основание в первом абзаце повторяется из письма в письмо."""
    text = ("Между ООО «Мосренстрой-6» и ООО «ФБ-Строй» заключен договор "
            "подряда №ЛС-СМР-Тайн от 18.05.2023 на выполнение работ")
    assert letters.find_contract(text) == "№ЛС-СМР-Тайн от 18.05.2023"
    assert letters.find_contract("договор поставки № 44/25 от 01.02.2025") \
        == "№44/25 от 01.02.2025"
    assert letters.find_contract("просто текст") is None


def test_search_key_holds_both_cases():
    """В письме «Вороновой», ищут «Воронова» — ключ должен покрыть оба."""
    key = letters.search_key("ООО «СК АВАНГАРД»", "Вороновой М.А.",
                             "Генеральному директору")
    assert "авангард" in key
    assert "вороновой" in key
    assert "воронов" in key          # начальная форма


def test_several_addressees_per_organization(env):
    conn, cfg = env
    conn.execute("UPDATE documents SET section = 'переписка'")
    rows = [
        {"org": "ООО «СК АВАНГАРД»", "person": "Вороновой М.А.",
         "position": "Генеральному директору", "letters": 7,
         "search_key": letters.search_key("ООО «СК АВАНГАРД»", "Вороновой М.А.")},
        {"org": "ООО «СК АВАНГАРД»", "person": "Петрову И.И.",
         "position": "Главному инженеру", "letters": 2,
         "search_key": letters.search_key("ООО «СК АВАНГАРД»", "Петрову И.И.")},
    ]
    db.save_contacts(conn, rows)
    saved = db.contacts(conn)
    assert len(saved) == 2
    assert saved[0]["person"] == "Вороновой М.А."      # чаще писали — выше


def test_intro_becomes_first_paragraph_without_slot(tmp_path: Path):
    """Если в бланке нет места под преамбулу, она идёт первым абзацем."""
    template = make_template(tmp_path / "бланк.docx")
    head = letters.Letterhead.from_config({**HEAD, "template": str(template)})
    lines = read_all(letters.render({
        "recipient_org": "ООО «Х»",
        "intro": "Между ООО «Х» и ООО «ФБ-СТРОЙ» заключен договор.",
        "body": "Просим ответить.",
    }, head))
    assert "Между ООО «Х» и ООО «ФБ-СТРОЙ» заключен договор." in lines
    assert "Просим ответить." in lines
    assert lines.index("Между ООО «Х» и ООО «ФБ-СТРОЙ» заключен договор.") \
        < lines.index("Просим ответить.")


def test_intro_uses_its_own_slot_when_present(tmp_path: Path):
    import docx

    path = tmp_path / "бланк.docx"
    document = docx.Document()
    document.add_paragraph("{{intro}}")
    document.add_paragraph("{{body}}")
    document.save(str(path))

    head = letters.Letterhead.from_config({**HEAD, "template": str(path)})
    lines = read_all(letters.render({
        "intro": "Основание письма.", "body": "Суть.",
    }, head))
    assert lines == ["Основание письма.", "Суть."]


def test_default_intro_has_the_usual_wording():
    head = letters.Letterhead.from_config({})
    assert "{{recipient_org}}" in head.intro
    assert "заключен договор" in head.intro


def test_intro_includes_object_and_contract():
    """Полная преамбула так, как её пишут в письмах."""
    filled = letters.fill_intro(letters.DEFAULT_INTRO, {
        "recipient_org": "ООО «Мосренстрой-6»",
        "own_org": "ООО «ФБ-Строй»",
        "contract": "№ЛС-СМР-Тайн от 18.05.2023",
        "object": "«Жилой дом по адресу: Тайнинская ул., вл.16, корп. 3»",
    })
    assert filled.startswith("Между ООО «Мосренстрой-6» и ООО «ФБ-Строй» "
                             "заключен договор подряда №ЛС-СМР-Тайн")
    assert "по объекту строительства: «Жилой дом" in filled
    assert filled.endswith("корп. 3».")


def test_intro_leaves_no_dangling_punctuation():
    """Незаполненный объект не должен оставить висящее двоеточие."""
    filled = letters.fill_intro(letters.DEFAULT_INTRO, {
        "recipient_org": "ООО «Х»", "own_org": "ООО «Y»"})
    assert filled.endswith("по объекту строительства.")
    assert ": ." not in filled
    assert "  " not in filled


def test_server_builds_intro_when_form_did_not(tmp_path: Path):
    template = make_template(tmp_path / "бланк.docx")
    head = letters.Letterhead.from_config({
        **HEAD, "template": str(template),
        "object": "«Жилой дом, Тайнинская ул., вл.16, корп. 3»",
    })
    lines = read_all(letters.render({
        "recipient_org": "ООО «Мосренстрой-6»",
        "contract": "№ЛС-СМР-Тайн от 18.05.2023",
        "body": "Суть письма.",
    }, head))
    assert any("заключен договор подряда №ЛС-СМР-Тайн" in line for line in lines)
    assert any("Тайнинская ул., вл.16, корп. 3" in line for line in lines)


def test_body_paragraphs_are_plain_by_default():
    assert letters.format_body("Первый.\nВторой.") == ["Первый.", "Второй."]


def test_body_can_be_numbered():
    assert letters.format_body("Первый.\nВторой.", "list") == \
        ["1. Первый.", "2. Второй."]


def test_empty_body():
    assert letters.format_body("") == [""]
    assert letters.format_body("", "list") == [""]


def test_intro_is_not_numbered(tmp_path: Path):
    """Преамбула — не пункт списка, нумерация начинается с текста."""
    template = make_template(tmp_path / "бланк.docx")
    head = letters.Letterhead.from_config({**HEAD, "template": str(template)})
    lines = read_all(letters.render({
        "recipient_org": "ООО «Х»",
        "intro": "Между сторонами заключен договор.",
        "body": "Первое требование.\nВторое требование.",
        "body_format": "list",
    }, head))
    assert "Между сторонами заключен договор." in lines
    assert "1. Первое требование." in lines
    assert "2. Второе требование." in lines


def test_numbering_is_limited_to_the_letters_folder(tmp_path: Path):
    """Номера живут в папке переписки: по всему архиву ловится посторонняя
    нумерация, из-за неё счётчик и улетал."""
    root = tmp_path / "arc"
    (root / "02 .Переписка (Письма)").mkdir(parents=True)
    (root / "Проекты").mkdir()
    (root / "02 .Переписка (Письма)" / "Исх РТП-175.txt").write_text(
        "Исх. РТП-175 от 19.03.2024 г.", encoding="utf-8")
    (root / "Проекты" / "РТП-9000.txt").write_text(
        "Исх. РТП-9000 от 01.01.2024 г.", encoding="utf-8")

    cfg = Config(roots=[Root(label="ПТО", path=str(root))],
                 db=str(tmp_path / "index.db"), letterhead=HEAD)
    conn = db.connect(cfg.db)
    try:
        indexer.run(conn, cfg)
        assert letters.next_number(conn, "РТП") == "РТП-9001"
        assert letters.next_number(conn, "РТП", "Переписка") == "РТП-176"
    finally:
        conn.close()


REAL_NAMES = [
    "Исх. РТП-1337 МРС (ответ УКМ-299).pdf",
    "Исх. РТП-1337 ПКС Инжиниринг.doc",     # то же письмо, исходник и скан
    "РТП-1336 МРС6.pdf",                    # без слова «Исх.»
    "Исх. РТП-1334  общее запрос акта сверки.pdf",
    "РТП-1331.pdf",
    "Приложение РТП-1327.pdf",
    "РТП-1317 КлассСтрой.docx",
]


def test_numbering_reads_real_file_names(tmp_path: Path):
    """Половина писем названа без «Исх.»: по одному разобранному полю
    такие номера не видны, а нумерация должна их учитывать."""
    box = tmp_path / "arc" / "02 .Переписка (Письма)"
    box.mkdir(parents=True)
    for name in REAL_NAMES:
        (box / name).write_bytes(b"")
    foreign = tmp_path / "arc" / "Проекты"
    foreign.mkdir()
    (foreign / "РТП-99000 посторонний.pdf").write_bytes(b"")

    cfg = Config(roots=[Root(label="ПТО", path=str(tmp_path / "arc"))],
                 db=str(tmp_path / "index.db"), letterhead=HEAD)
    conn = db.connect(cfg.db)
    try:
        indexer.run(conn, cfg)
        assert letters.next_number(conn, "РТП", "Переписка") == "РТП-1338"
        # без ограничения папкой в счёт идёт посторонняя нумерация
        assert letters.next_number(conn, "РТП") == "РТП-99001"
    finally:
        conn.close()


def test_number_from_file_name_without_prefix_word():
    from docsearch import meta

    assert meta.find_number("РТП-1336 МРС6", from_start=True) == "РТП-1336"
    assert meta.find_number("РТП-1331", from_start=True) == "РТП-1331"
    assert meta.find_number("Исх. РТП-1337 МРС", from_start=True) == "РТП-1337"


def test_object_code_is_not_mistaken_for_a_letter_number():
    """Шифр альбома начинается с цифр и номером письма быть не может."""
    from docsearch import meta

    assert meta.find_number("208-1121-ОК-1-АР4", from_start=True) is None


def test_template_builder_strips_word_numbering(tmp_path: Path):
    """Нумерация из исходного письма рисовала «1.» и «2.» на пустых абзацах."""
    import importlib.util

    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    source = tmp_path / "письмо.docx"
    document = docx.Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].text = "Исх. РТП-1 от 01.01.2026"
    document.add_paragraph("Тема: «Проверка»")
    document.add_paragraph("Уважаемый Иван Иванович!")
    document.add_paragraph("Между сторонами заключен договор.")
    document.add_paragraph("Текст письма.")
    # нумерация бывает прямой и унаследованной от стиля — проверяем обе
    direct = document.add_paragraph("")
    pPr = direct._element.get_or_add_pPr()
    pPr.append(OxmlElement("w:numPr"))
    styled = document.add_paragraph("", style="List Number")
    document.add_paragraph("С уважением,")
    document.add_paragraph("Руководитель" + chr(9) * 2 + "Иванов И.И.")
    assert pPr.find(qn("w:numPr")) is not None
    assert "List" in styled.style.name
    document.save(str(source))

    spec = importlib.util.spec_from_file_location(
        "make_template",
        Path(__file__).resolve().parent.parent / "tools" / "make_template.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    target = tmp_path / "шаблон.docx"
    module.build(str(source), str(target))

    result = docx.Document(str(target))
    for paragraph in result.paragraphs:
        pPr = paragraph._element.find(qn("w:pPr"))
        if pPr is not None:
            assert pPr.find(qn("w:numPr")) is None
        if not paragraph.text.strip():
            assert "List" not in (paragraph.style.name or "")


def test_own_organization_printed_in_full(tmp_path: Path):
    """В преамбуле должно стоять печатное название, а не строка сопоставления."""
    from docsearch.config import Config, Root
    from docsearch.web import create_app

    cfg = Config(roots=[Root(label="П", path=str(tmp_path))],
                 db=str(tmp_path / "index.db"),
                 own_org="ФБ-Строй", letterhead=HEAD)
    conn = db.connect(cfg.db)
    conn.close()
    client = TestClient(create_app(cfg))
    assert client.get("/api/letter/draft").json()["own_org"] == "ООО «ФБ-СТРОЙ»"


def test_contract_attached_without_parsed_addressee(env):
    """Оформление бланков разное: основание есть даже там, где блок
    «кому» не разобрался."""
    conn, cfg = env
    doc_id = conn.execute("SELECT MIN(id) id FROM documents").fetchone()["id"]
    conn.execute(
        "UPDATE documents SET section='переписка', doc_type='письмо',"
        " counterparty='ООО «НЛ Групп»' WHERE id=?", (doc_id,))
    body = chr(10).join([
        "Генеральному директору",
        "ООО «НЛ Групп»",
        "Иванову И.И.",
        "Между ООО «НЛ Групп» и ООО «ФБ-Строй» заключен договор подряда"
        " №НЛ-17 от 03.04.2024 на выполнение работ",
    ])
    conn.execute("DELETE FROM doc_fts WHERE rowid = ?", (doc_id,))
    conn.execute("INSERT INTO doc_fts (rowid, name, body, lemmas)"
                 " VALUES (?,?,?,?)", (doc_id, "письмо", body, body))
    conn.commit()

    rows = letters.rebuild_contacts(conn)
    found = {r["org"]: r for r in rows}
    assert found["ООО «НЛ Групп»"]["contract"] == "№НЛ-17 от 03.04.2024"


def _letter_body(position, org, person, contract=None):
    lines = [position, org, person]
    if contract:
        lines.append(f"Между {org} и ООО «ФБ-Строй» заключен договор подряда"
                     f" {contract} на выполнение работ")
    return chr(10).join(lines)


def _put(conn, doc_id, body, counterparty=None):
    conn.execute("UPDATE documents SET section='переписка', doc_type='письмо',"
                 " counterparty=? WHERE id=?", (counterparty, doc_id))
    conn.execute("DELETE FROM doc_fts WHERE rowid = ?", (doc_id,))
    conn.execute("INSERT INTO doc_fts (rowid, name, body, lemmas)"
                 " VALUES (?,?,?,?)", (doc_id, "письмо", body, body))


def test_contacts_merge_spelling_variants(env):
    """В справочнике одна организация должна быть одной строкой."""
    conn, cfg = env
    ids = [r["id"] for r in conn.execute("SELECT id FROM documents LIMIT 2")]
    _put(conn, ids[0], _letter_body("Генеральному директору",
                                    "ООО «ПД-Проект»", "Демину А.А."))
    _put(conn, ids[1], _letter_body("Генеральному директору",
                                    "ООО «ПД-ПРОЕКТ»", "Демин А.А."))
    conn.commit()

    rows = letters.rebuild_contacts(conn)
    assert len(rows) == 1
    assert rows[0]["letters"] == 2


def test_own_organization_is_not_an_addressee(env):
    """Входящие письма адресованы нам — в справочнике адресатов нам не место."""
    conn, cfg = env
    ids = [r["id"] for r in conn.execute("SELECT id FROM documents LIMIT 2")]
    _put(conn, ids[0], _letter_body("Генеральному директору",
                                    "ООО «ФБ-Строй»", "Глухову А.В."))
    _put(conn, ids[1], _letter_body("Генеральному директору",
                                    "ООО «НЛ-ГРУПП»", "Пану В.А."))
    conn.commit()

    rows = letters.rebuild_contacts(conn, own_org="ФБ-Строй")
    orgs = [r["org"] for r in rows]
    assert "ООО «НЛ-ГРУПП»" in orgs
    assert not any("ФБ" in org for org in orgs)


def test_contract_shared_across_spellings(env):
    """Договор найден в письме с одним написанием — нужен и при другом."""
    conn, cfg = env
    ids = [r["id"] for r in conn.execute("SELECT id FROM documents LIMIT 2")]
    _put(conn, ids[0], _letter_body("Генеральному директору",
                                    "ООО «СК АВАНГАРД»", "Вороновой М.А.",
                                    "№СКА-5 от 10.01.2025"))
    _put(conn, ids[1], _letter_body("Генеральному директору",
                                    "ООО «СК-Авангард»", "Воронова М.А."))
    conn.commit()

    rows = letters.rebuild_contacts(conn)
    assert len(rows) == 1
    assert rows[0]["contract"] == "№СКА-5 от 10.01.2025"


def test_contract_goes_to_the_party_not_the_addressee():
    """В преамбуле названы стороны договора, а письмо о нём может уйти
    и третьему лицу: у ПД-Проекта не наш договор с Мосренстроем."""
    text = chr(10).join([
        "Генеральному директору", "ООО «ПД-Проект»", "Демину А.А.",
        "Между ООО «Мосренстрой-6» и ООО «ФБ-Строй» заключен договор подряда"
        " №ЛС-СМР-Тайн от 18.05.2023 на выполнение работ",
    ])
    party, contract = letters.find_contract_parties(text, "ФБ-Строй")
    assert party == "ООО «Мосренстрой-6»"
    assert contract == "№ЛС-СМР-Тайн от 18.05.2023"


def test_contract_parties_skip_our_own_side():
    party, _ = letters.find_contract_parties(
        "Между ООО «НЛ-ГРУПП» и ООО «ФБ-Строй» заключен договор"
        " №РЕН-2804 от 28.04.2025", "ФБ-Строй")
    assert party == "ООО «НЛ-ГРУПП»"


def test_nameless_entry_merges_into_named_one():
    """«Генеральному директору ООО «X»» без фамилии — тот же адресат."""
    rows = [
        {"org": "ООО «ПД-Проект»", "person": "", "position": "Генеральному директору",
         "contract": "№5 от 01.01.2025", "letters": 29, "last_date": None},
        {"org": "ООО «ПД-ПРОЕКТ»", "person": "Демину А.А.",
         "position": "Генеральному директору", "contract": None,
         "letters": 2, "last_date": None},
    ]
    merged = letters.merge_nameless(rows)
    assert len(merged) == 1
    assert merged[0]["person"] == "Демину А.А."
    assert merged[0]["letters"] == 31
    assert merged[0]["contract"] == "№5 от 01.01.2025"


def test_nameless_entry_survives_without_a_named_one():
    """У М-СТРОЙ фамилии нет нигде — строка должна остаться."""
    rows = [{"org": "ООО «М-СТРОЙ»", "person": "",
             "position": "Генеральному директору", "contract": None,
             "letters": 9, "last_date": None}]
    assert len(letters.merge_nameless(rows)) == 1


def test_different_positions_stay_separate():
    rows = [
        {"org": "ООО «Мосренстрой-6»", "person": "Горбуновой М.Н.",
         "position": "Заместителю генерального", "contract": None,
         "letters": 90, "last_date": None},
        {"org": "ООО «Мосренстрой-6»", "person": "",
         "position": "Генеральному директору", "contract": None,
         "letters": 5, "last_date": None},
    ]
    assert len(letters.merge_nameless(rows)) == 2


def test_stray_numbering_cleaned_at_render(tmp_path: Path):
    """Старый бланк со списком не должен рисовать «1.» в новом письме."""
    import docx
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    path = tmp_path / "бланк.docx"
    document = docx.Document()
    document.add_paragraph("{{body}}")
    stray = document.add_paragraph("")
    stray._element.get_or_add_pPr().append(OxmlElement("w:numPr"))
    document.add_paragraph("", style="List Number")
    document.save(str(path))

    head = letters.Letterhead.from_config({**HEAD, "template": str(path)})
    blob = letters.render({"recipient_org": "ООО «Х»", "body": "Текст."}, head)

    result = docx.Document(BytesIO(blob))
    for paragraph in result.paragraphs:
        if paragraph.text.strip():
            continue
        pPr = paragraph._element.find(qn("w:pPr"))
        if pPr is not None:
            assert pPr.find(qn("w:numPr")) is None
        assert "List" not in (paragraph.style.name or "")


INCOMING = chr(10).join([
    "№ 114-25/ИСХ от 19.11.25",
    "Генеральному директору",
    "ГУП «МОСВОДОСТОК»",
    "Ишханяну К.Р.",
    "Уважаемый Константин Рафаэлович!",
    "Общество с ограниченной ответственностью «ЕРЛУК» выполняет работы"
    " по объекту строительства.",
    "Генеральный директор А.М. Лукашин",
])


def test_sender_parsed_from_signature():
    """Входящее письмо даёт контрагента, которому мы не писали."""
    sender = letters.parse_sender(INCOMING, "ФБ-Строй", "ГУП «МОСВОДОСТОК»")
    assert sender["org"] == "ООО «ЕРЛУК»"
    assert sender["person"] == "Лукашину А.М."
    assert sender["position"] == "Генеральному директору"


def test_sender_is_converted_to_addressee_form():
    """В подписи именительный падеж и инициалы впереди — в адресате наоборот."""
    person, position = letters.to_addressee_form("А.М. Лукашин",
                                                 "Генеральный директор")
    assert person == "Лукашину А.М."
    assert position == "Генеральному директору"

    person, _ = letters.to_addressee_form("Воронова М.А.", "Директор")
    assert person == "Вороновой М.А."


def test_position_tail_is_not_inflected():
    """«Руководитель проекта»: склоняется только главное слово."""
    _, position = letters.to_addressee_form("А.А. Иванов", "Руководитель проекта")
    assert position == "Руководителю проекта"


def test_sender_skips_us_and_the_addressee():
    assert letters.parse_sender(INCOMING, "ЕРЛУК", "ГУП «МОСВОДОСТОК»") is None


def test_letter_without_signature_has_no_sender():
    assert letters.parse_sender("Просто текст без подписи", "ФБ-Строй") is None


def test_incoming_letter_adds_its_sender(env):
    conn, cfg = env
    doc_id = conn.execute("SELECT MIN(id) id FROM documents").fetchone()["id"]
    _put(conn, doc_id, INCOMING)
    conn.commit()

    rows = letters.rebuild_contacts(conn, own_org="ФБ-Строй")
    orgs = {r["org"]: r for r in rows}
    assert "ООО «ЕРЛУК»" in orgs
    assert orgs["ООО «ЕРЛУК»"]["person"] == "Лукашину А.М."
    assert "ГУП «МОСВОДОСТОК»" in orgs       # адресат тоже на месте


def test_greeting_parsed_from_letter():
    """Имя и отчество получателя есть только в обращении."""
    text = chr(10).join([
        "Генеральному директору", "ГУП «МОСВОДОСТОК»", "Ишханяну К.Р.",
        "Уважаемый Константин Рафаэлович !", "Текст письма.",
    ])
    assert letters.parse_greeting(text) == "Уважаемый Константин Рафаэлович!"


def test_greeting_keeps_gender():
    assert letters.parse_greeting("Уважаемая Марина Александровна") \
        == "Уважаемая Марина Александровна!"


def test_no_greeting():
    assert letters.parse_greeting("Письмо без обращения") is None
    assert letters.parse_greeting("") is None


def test_greeting_saved_with_the_contact(env):
    conn, cfg = env
    doc_id = conn.execute("SELECT MIN(id) id FROM documents").fetchone()["id"]
    _put(conn, doc_id, chr(10).join([
        "Генеральному директору", "ООО «СК АВАНГАРД»", "Вороновой М.А.",
        "Уважаемая Марина Александровна!", "Текст.",
    ]))
    conn.commit()

    rows = letters.rebuild_contacts(conn, own_org="ФБ-Строй")
    found = {r["org"]: r for r in rows}
    assert found["ООО «СК АВАНГАРД»"]["greeting"] == "Уважаемая Марина Александровна!"


def test_blank_paragraphs_collapse(tmp_path: Path):
    """Вереницу пустых абзацев из бланка приходилось вычищать руками."""
    import docx

    path = tmp_path / "бланк.docx"
    document = docx.Document()
    document.add_paragraph("{{body}}")
    for _ in range(5):
        document.add_paragraph("")
    document.add_paragraph("С уважением,")
    document.save(str(path))

    head = letters.Letterhead.from_config({**HEAD, "template": str(path)})
    blob = letters.render({"recipient_org": "ООО «Х»", "body": "Текст."}, head)

    result = docx.Document(BytesIO(blob))
    runs = 0
    longest = 0
    for paragraph in result.paragraphs:
        runs = 0 if paragraph.text.strip() else runs + 1
        longest = max(longest, runs)
    assert longest <= 1
