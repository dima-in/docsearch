"""Акты комиссии: бланк, номер по папке, комиссия из прошлых актов.

В архиве у актов один бланк на все случаи: шапка с объектом, номер и
дата, «Комиссия в составе», «составили настоящий Акт в том, что…» и
подписи. Меняется предмет акта — осмотр работ по КС-2, повреждение,
передача помещений, — а остальное от акта к акту одинаковое, и всё это
до сих пор набирали руками.

Комиссию в бланке пишут дважды: вверху и в подписях. Здесь она задаётся
один раз списком и сама попадает в оба места. Справочник людей собран
из тех же актов: должность, организация и фамилия — как их писали.

Выравнивание в исходном бланке сделано пробелами и подчёркиваниями, и
стоит фамилии оказаться длиннее, строка съезжает. В шаблоне вместо них
табуляция: дата прижата к краю, линия подписи на своём месте при любой
длине фамилии.
"""
from __future__ import annotations

import copy
import io
import re
import sqlite3
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn
from docx.shared import Twips
from docx.text.paragraph import Paragraph

from . import letters

ABSENT = "отсутствовал, акт составлен в одностороннем порядке"
BODY_LEAD = "составили настоящий Акт в том, что "
DEFAULT_FOLDER = "АКТЫ"

# — поля пакета: {шифр}, {КС-2} — в одинарных скобках, чтобы не путать
# с разметкой шаблона {{number}}
RE_FIELD = re.compile(r"(?<!\{)\{([^{}\n]{1,40})\}(?!\})")
# — стороны: [кто передаёт], [кто принимает] — организации, которые в форме
# выбираются из списка. Сколько разных скобок, столько и списков
RE_PARTY = re.compile(r"\[([А-ЯЁа-яё][А-ЯЁа-яё -]{1,38})\]")


@dataclass
class ActSettings:
    template: str = ""
    folder: str = DEFAULT_FOLDER
    organization: str = ""
    presets: list = field(default_factory=list)

    @classmethod
    def from_config(cls, cfg) -> "ActSettings":
        raw = getattr(cfg, "acts", None) or {}
        organization = (raw.get("organization") or cfg.own_org or "").strip()
        return cls(
            template=raw.get("template") or "",
            folder=(raw.get("folder") or DEFAULT_FOLDER).strip(),
            organization=organization,
            presets=list(raw.get("presets") or []) or default_presets(),
        )

    @property
    def ready(self) -> bool:
        return bool(self.template) and Path(self.template).exists()


def default_presets() -> list[dict]:
    """Заготовки текста — по самым частым актам из архива.

    Прочерки — места, которые всё равно придётся вписать. Стороны в
    квадратных скобках выбираются в форме из списка организаций. Поля в
    фигурных скобках — колонки пакета: на каждую строку выйдет отдельный акт.
    """
    return [
        {"name": "Свободный текст", "text": ""},
        {"name": "Осмотр работ по КС-2 (можно пакетом)",
         "text": "произведен осмотр работ, заявленных в КС-2 за номером {КС-2} "
                 "от {дата КС-2} шифр {шифр}.\n"
                 "СМР по данному разделу не приняты. {замечания}\n"
                 "Не представлена исполнительная документация, "
                 "предусмотренная Договором."},
        {"name": "Повреждение работ",
         "text": "во время производства работ подрядной организацией [кто повредил] "
                 "были повреждены работы, выполненные [чьи работы].\n"
                 "Объём восстановительных работ составляет:\n- ___\n"
                 "Комиссией принято решение: ___"},
        {"name": "Передача помещений под работы",
         "text": "представителями [кто передаёт] передаются помещения на ___ секции "
                 "___ этаж для производства работ по ___, представителями "
                 "[кто принимает] выполнена приёмка вышеуказанных помещений.\n"
                 "Дополнения: ___"},
        {"name": "Демонтаж с сохранением материала",
         "text": "согласно ___ силами [кто демонтировал] была демонтирована ___ "
                 "в объёме ___ тн, в том числе с сохранением материала:\n"
                 "- ___ – ___ тн"},
        {"name": "Списание",
         "text": "___ имеет износ и дефекты, появившиеся в процессе длительной "
                 "эксплуатации. Ремонт нецелесообразен.\n"
                 "В связи с вышесказанным принято решение о списании и утилизации."},
    ]


def full_org(name: str, people: list[dict]) -> str:
    """Название с формой собственности, как оно пишется в актах.

    В конфиге своя организация часто записана коротко — «ФБ-Строй», а в
    шапке и тексте акта нужно «ООО «ФБ-Строй»». Форму берём из прошлых
    актов той же организации; не нашли — оставляем как есть, не гадаем.
    """
    name = (name or "").strip()
    if not name or RE_ORG.fullmatch(name):
        return name
    key = letters.org_key(name)
    for person in people:
        org = (person.get("org") or "").strip()
        if letters.org_key(org) == key and RE_ORG.fullmatch(org):
            return org
    return name


# — член комиссии ———————————————————————————————————————————————

FORMS = r"(?:ООО|АО|ПАО|ЗАО|ОАО|НАО|ГБУ|ГКУ|ГУП|МУП|ФГУП|ФГБУ|АНО)"
RE_ORG = re.compile(
    FORMS + r"\s*[«\"“][^»\"”]{1,60}[»\"”]"
    r"|ИП\s*«[^»]{1,40}»"
    r"|ИП\s+[А-ЯЁ][а-яё-]+(?:\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.?)?"
)
RE_NAME = re.compile(
    r"^[\s,.–—»-]*([А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?)\s+([А-ЯЁ])\.\s*([А-ЯЁ])\.?"
)
# дальше комиссии не читаем: в тексте акта организации тоже упоминаются
RE_COMMISSION = re.compile(r"комиссия\s+в\s+составе\s*:?", re.IGNORECASE)
RE_COMMISSION_END = re.compile(r"^\s*составил", re.IGNORECASE | re.MULTILINE)


@dataclass
class Member:
    position: str = ""
    org: str = ""
    name: str = ""
    absent: bool = False

    @classmethod
    def from_dict(cls, data: dict) -> "Member":
        return cls(
            position=_tidy(data.get("position")),
            org=_tidy(data.get("org")),
            name=_tidy(data.get("name")),
            absent=bool(data.get("absent")),
        )

    def as_dict(self) -> dict:
        return {"position": self.position, "org": self.org, "name": self.name,
                "absent": self.absent}

    def commission_line(self) -> str:
        """Строка в «Комиссия в составе». Табуляция в конце — линия до края."""
        if self.absent:
            position = self.position or "Представитель"
            return " ".join(filter(None, [position, self.org, ABSENT]))
        return " ".join(filter(None, [self.position, self.org, self.name])) + "\t"

    def signature(self) -> dict:
        """Подпись: должность; организация — линия — фамилия; «(подпись)».

        Длинное название организации не помещается до линии подписи —
        тогда оно уходит строкой к должности.
        """
        if self.absent:
            return {"position": self.position or "Представитель",
                    "sign": f"{self.org} {ABSENT}".strip(), "caption": ""}
        name = self.name or "_" * 14
        if len(self.org) > LONG_ORG:
            position = "\n".join(filter(None, [self.position, self.org]))
            return {"position": position, "sign": f"\t\t\t{name}",
                    "caption": "\t(подпись)"}
        return {"position": self.position, "sign": f"{self.org}\t\t\t{name}",
                "caption": "\t(подпись)"}


# длиннее жирным шрифтом не помещается до линии подписи
LONG_ORG = 24


def _tidy(value) -> str:
    return " ".join(str(value or "").split())


def _name(surname: str, first: str, middle: str) -> str:
    return f"{surname} {first}.{middle}."


def parse_member(line: str) -> Member | None:
    """«Руководитель проекта ООО «Альфа» Петров И.И.____» -> член комиссии."""
    text = _tidy(line.replace("_", " ")).strip(" .")
    if not text or text.startswith("("):
        return None
    match = RE_ORG.search(text)
    if not match:
        return None
    # «ООО«Альфа»» без пробела — так набрано в части актов
    org = re.sub(r"^(" + FORMS + r"|ИП)\s*([«\"“])", r"\1 \2", _tidy(match.group(0)))
    if org.startswith("ИП"):
        found = RE_NAME.match(org[2:])
        if found:
            org = "ИП " + _name(*found.groups())
    position = text[:match.start()].strip(" ,.–—-")
    rest = text[match.end():]
    absent = "отсутств" in rest.lower()
    name = ""
    if not absent and not org.startswith("ИП"):
        found = RE_NAME.match(rest)
        if found:
            name = _name(*found.groups())
    return Member(position=position, org=org, name=name, absent=absent)


def commission_of(text: str) -> list[Member]:
    """Комиссия из текста акта — между «Комиссия в составе» и «составили»."""
    start = RE_COMMISSION.search(text or "")
    if not start:
        return []
    rest = text[start.end():]
    end = RE_COMMISSION_END.search(rest)
    block = rest[:end.start()] if end else rest[:1500]
    members = []
    for line in block.splitlines():
        member = parse_member(line)
        if member:
            members.append(member)
    return members


# — папка актов: номера и люди ————————————————————————————————————

RE_ACT_NUMBER = re.compile(r"\bакт\s*№\s*(\d{1,4})", re.IGNORECASE)
HEAD_CHARS = 600            # номер акта стоит в шапке


def _in_folder(folder: str) -> tuple[str, list]:
    """Только сама папка актов: «Акты ЗОС» и «Акты полигон» — другие серии."""
    prefix = folder.strip().strip("\\/").lower()
    if not prefix:
        return "1", []
    return ("(ru_lower(d.rel_path) LIKE ? OR ru_lower(d.rel_path) LIKE ?)",
            [prefix + chr(92) + "%", prefix + "/%"])


def _rows(conn: sqlite3.Connection, folder: str):
    where, params = _in_folder(folder)
    return conn.execute(
        "SELECT d.name, d.rel_path, d.doc_date, d.mtime, f.body AS body "
        "FROM documents d LEFT JOIN doc_fts f ON f.rowid = d.id "
        f"WHERE {where}", params).fetchall()


def act_number(name: str, body: str) -> int | None:
    for source in (name or "", (body or "")[:HEAD_CHARS]):
        found = RE_ACT_NUMBER.search(source)
        if found:
            return int(found.group(1))
    return None


def numbers(conn: sqlite3.Connection, folder: str = DEFAULT_FOLDER) -> list[int]:
    found = [n for row in _rows(conn, folder)
             if (n := act_number(row["name"], row["body"])) is not None]
    return sorted(found)


def next_number(conn: sqlite3.Connection, folder: str = DEFAULT_FOLDER) -> int:
    """Следующий номер. Как с письмами: журнал — это сама папка."""
    return letters.highest_sane(numbers(conn, folder)) + 1


def _when(row) -> str:
    if row["doc_date"]:
        return row["doc_date"]
    return datetime.fromtimestamp(row["mtime"]).date().isoformat()


def directory(conn: sqlite3.Connection, folder: str = DEFAULT_FOLDER,
              organization: str = "") -> tuple[list[dict], dict | None]:
    """Справочник членов комиссии и состав последнего акта.

    Один человек пишется по-разному: «Д. А.» и «Д.А.», «СК-АВАНГАРД» и
    «СК-Авангард». Сводим по ключам организации и фамилии, а показываем
    самое свежее написание.
    """
    ours = letters.org_key(organization) if organization else ""
    people: dict[tuple, dict] = {}
    acts_found: list[dict] = []

    for row in sorted(_rows(conn, folder), key=_when):
        members = commission_of(row["body"] or "")
        if not members:
            continue
        when = _when(row)
        number = act_number(row["name"], row["body"])
        acts_found.append({"number": number, "date": when,
                           "members": [m.as_dict() for m in members]})
        for member in members:
            org = letters.org_key(member.org)
            who = (letters.person_key(member.name) if member.name
                   else "?" + letters.position_key(member.position))
            entry = people.setdefault((org, who), {"count": 0})
            entry.update(position=member.position or entry.get("position", ""),
                         org=member.org, name=member.name, last=when,
                         own=bool(ours) and org == ours)
            entry["count"] += 1

    # «Представитель ООО X» без фамилии нужен, только пока фамилий от X нет
    named = {org for (org, who), e in people.items() if e["name"]}
    rows = [e for (org, who), e in people.items() if e["name"] or org not in named]
    for entry in rows:
        entry["search_key"] = letters.search_key(entry["position"], entry["org"],
                                                 entry["name"])
    # свежие вперёд, а поверх — наши и частые: сортировка устойчивая
    rows.sort(key=lambda e: e["last"], reverse=True)
    rows.sort(key=lambda e: (not e["own"], -e["count"]))
    return rows, _last_act(acts_found)


def _last_act(found: list[dict]) -> dict | None:
    """Последний акт серии — с наибольшим номером, а не с поздней датой.

    Дату индекс берёт из текста, и в акте с номером поменьше может
    упоминаться дата позже: «КС-2 от 15.09» в акте, составленном раньше.
    Нумерация сквозная, ей и верим. Одинокий выброс отбрасываем так же,
    как при выборе следующего номера.
    """
    if not found:
        return None
    numbered = [a for a in found if a["number"] is not None]
    if numbered:
        top = letters.highest_sane(sorted(a["number"] for a in numbered))
        same = [a for a in numbered if a["number"] == top]
        if same:
            return max(same, key=lambda a: a["date"])
    return max(found, key=lambda a: a["date"])


# — шаблон ——————————————————————————————————————————————————————

PLACEHOLDER = re.compile(r"{{\s*([?#/]?)(\w+)\s*}}")


def _paragraph(element, document) -> Paragraph:
    return Paragraph(element, document._body)


def _all_paragraphs(document):
    """Абзацы в порядке документа, включая ячейки: бланк свёрстан таблицей."""
    for element in document.element.body.iter(qn("w:p")):
        yield _paragraph(element, document)


def _text_width(element, document) -> int:
    """Ширина строки в твипах: ячейки, если абзац в таблице, иначе листа."""
    cell = next((a for a in element.iterancestors(qn("w:tc"))), None)
    if cell is not None:
        width = cell.find(f"{qn('w:tcPr')}/{qn('w:tcW')}")
        if width is not None and width.get(qn("w:type")) in (None, "dxa"):
            return int(width.get(qn("w:w")))
    section = document.sections[0]
    return int((section.page_width - section.left_margin - section.right_margin)
               / 635)     # EMU в твипе


def _tabs(paragraph, stops: list[tuple]) -> None:
    """Заменить позиции табуляции: (твипы, выравнивание, заполнитель)."""
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.clear_all()
    for position, alignment, leader in stops:
        tab_stops.add_tab_stop(Twips(position), alignment, leader)


def _clone_after(anchor, text: str, document) -> Paragraph:
    clone = copy.deepcopy(anchor._element)
    anchor._element.addnext(clone)
    paragraph = _paragraph(clone, document)
    letters._set_text(paragraph, text)
    return paragraph


def _replace_blank(paragraph, value: str) -> None:
    """«АКТ № _____» -> «АКТ № {{number}}»: прочерк меняем, текст до него
    оставляем.

    Прочерк в бланке бывает набран несколькими прогонами, и оформление у
    них разное: «Строительная организация ____» — первый прочерк простой,
    а дальше жирный с подчёркиванием, и в заполненных актах название
    стоит именно так. Значение ставим в первый выделенный прогон.
    """
    blanks = [run for run in paragraph.runs if "_" in run.text]
    if not blanks:
        paragraph.add_run(" " + value)
        return
    chosen = next((run for run in blanks if run.underline or run.bold), blanks[0])
    for run in blanks:
        if run is chosen:
            run.text = run.text.split("_")[0] + value
        else:
            run.text = run.text.replace("_", "")


def _remove(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def _is_dashes(text: str) -> bool:
    return "_" in text and not re.sub(r"[_\s.,]", "", text)


def make_template(source: str | Path):
    """Шаблон из вашего бланка акта: размеченный .docx.

    Ищем в бланке опорные строки — «Строительная организация», «АКТ №»,
    «г. Москва», «Комиссия в составе», «составили», «(Основание» — и
    меняем вокруг них только то, что заполняется. Остальное — шапка с
    объектом, шрифты, поля — остаётся вашим.
    """
    document = Document(str(source))
    paragraphs = list(_all_paragraphs(document))

    def first(predicate, after=None):
        start = paragraphs.index(after) + 1 if after is not None else 0
        return next((p for p in paragraphs[start:] if predicate(p.text)), None)

    org = first(lambda t: t.strip().lower().startswith("строительная организация"))
    if org is not None:
        _replace_blank(org, "{{organization}}")

    number = first(lambda t: re.match(r"\s*акт\s*№", t, re.IGNORECASE))
    if number is None:
        raise ValueError("В бланке нет строки «АКТ №» — это точно бланк акта?")
    _replace_blank(number, "{{number}}")

    when = first(lambda t: t.strip().startswith("г.") and "«" in t, after=number)
    if when is not None:
        city = when.text.split("«")[0].strip()
        letters._set_text(when, f"{city}\t« {{{{day}}}} » {{{{month}}}} {{{{year}}}} г.")
        width = _text_width(when._element, document)
        _tabs(when, [(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.SPACES)])

    commission = first(lambda t: RE_COMMISSION.search(t))
    body = first(lambda t: t.strip().lower().startswith("составил"),
                 after=commission)
    if commission is None or body is None:
        raise ValueError("В бланке нет «Комиссия в составе» или «составили "
                         "настоящий Акт» — разметить нечего")
    width = _text_width(body._element, document)

    # — комиссия: одна пара «строка — подпись под ней», дальше повтор
    between = paragraphs[paragraphs.index(commission) + 1:paragraphs.index(body)]
    line_proto = next((p for p in between if "_" in p.text), None)
    caption_proto = next((p for p in between if p.text.strip().startswith("(")), None)
    anchor = commission
    anchor = _clone_after(anchor, "{{#commission}}", document)
    if line_proto is not None:
        line = copy.deepcopy(line_proto._element)
        anchor._element.addnext(line)
        anchor = _paragraph(line, document)
        letters._set_text(anchor, "{{line}}")
        _tabs(anchor, [(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.LINES)])
    else:
        anchor = _clone_after(anchor, "{{line}}", document)
    if caption_proto is not None:
        caption = copy.deepcopy(caption_proto._element)
        anchor._element.addnext(caption)
        anchor = _paragraph(caption, document)
    anchor = _clone_after(anchor, "{{/commission}}", document)
    for paragraph in between:
        _remove(paragraph)
    _clone_after(anchor, "", document)

    # — текст акта
    letters._set_text(body, BODY_LEAD + "{{body}}")
    after_body = paragraphs[paragraphs.index(body) + 1:]
    while after_body and _is_dashes(after_body[0].text):
        _remove(after_body.pop(0))
    appendix = _clone_after(body, "Приложение: {{?appendix}}", document)

    basis = first(lambda t: t.strip().startswith("(Основание"), after=body)
    if basis is not None:
        letters._set_text(basis, "(Основание: {{?basis}})")
        tail_from = basis
    else:
        tail_from = appendix

    # — подписи: всё, что в бланке после текста акта, уходит, вместо него
    # повтор. Основание остаётся: это часть текста, а не подписи
    container = body._element.getparent()
    tail = [p for p in paragraphs[paragraphs.index(body) + 1:]
            if p is not basis and p._element.getparent() is container]
    sign_proto = next((p for p in tail if "_" in p.text), None) or body
    caption_proto = next((p for p in tail if "(подпись)" in p.text), None) or body
    empty_proto = next((p for p in tail if not p.text.strip()), None) or body

    # как в ваших актах: линия подписи с середины строки, фамилия с отступом
    start, end, name_at = (int(width * k) for k in (0.40, 0.66, 0.70))
    anchor = tail_from
    for _ in range(2):
        anchor = _clone_after_proto(anchor, empty_proto, "", document)
    anchor = _clone_after(anchor, "{{#signatures}}", document)
    anchor = _clone_after_proto(anchor, sign_proto, "{{position}}", document)
    _tabs(anchor, [])
    anchor = _clone_after_proto(anchor, sign_proto, "{{sign}}", document)
    _tabs(anchor, [(start, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES),
                   (end, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.LINES),
                   (name_at, WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.SPACES)])
    anchor = _clone_after_proto(anchor, caption_proto, "{{caption}}", document)
    _tabs(anchor, [((start + end) // 2, WD_TAB_ALIGNMENT.CENTER,
                    WD_TAB_LEADER.SPACES)])
    anchor = _clone_after_proto(anchor, empty_proto, "", document)
    _clone_after(anchor, "{{/signatures}}", document)
    for paragraph in tail:
        _remove(paragraph)
    return document


def _clone_after_proto(anchor, proto, text: str, document) -> Paragraph:
    clone = copy.deepcopy(proto._element)
    anchor._element.addnext(clone)
    paragraph = _paragraph(clone, document)
    letters._set_text(paragraph, text)
    paragraph.paragraph_format.alignment = proto.paragraph_format.alignment
    return paragraph


def placeholders(document) -> list[str]:
    found = []
    for paragraph in _all_paragraphs(document):
        for flag, name in PLACEHOLDER.findall(paragraph.text):
            token = f"{flag}{name}"
            if token not in found:
                found.append(token)
    return found


# — заполнение ——————————————————————————————————————————————————

def _fill(paragraph, context: dict, document) -> bool:
    """Подставить значения. False — абзац надо убрать.

    Убираем абзац, где пусто необязательное поле {{?x}}, и абзац, который
    после подстановки опустел совсем: незачем оставлять дыру на месте
    «(подпись)» у отсутствовавшего.

    Меняем внутри прогона, где стоит подстановка: у организации в бланке
    своё оформление — жирный с подчёркиванием, — и переписывать всё в
    первый прогон значит его потерять. Многострочное значение даёт новые
    абзацы с оформлением исходного.
    """
    text = paragraph.text
    if "{{" not in text:
        return True
    for flag, name in PLACEHOLDER.findall(text):
        if flag == "?" and not str(context.get(name) or "").strip():
            return False
    if not PLACEHOLDER.sub(lambda m: str(context.get(m.group(2)) or ""), text).strip():
        return False

    rest: list[str] = []

    def first_line(match) -> str:
        head, *tail = str(context.get(match.group(2)) or "").split("\n")
        rest.extend(tail)
        return head

    runs = paragraph.runs
    whole = len(PLACEHOLDER.findall(text))
    if whole == sum(len(PLACEHOLDER.findall(run.text)) for run in runs):
        for run in runs:
            if "{{" in run.text:
                run.text = PLACEHOLDER.sub(first_line, run.text)
    else:
        # подстановку разрезало между прогонами — правили шаблон в Word
        letters._set_text(paragraph, PLACEHOLDER.sub(first_line, text))

    anchor = paragraph._element
    for line in rest:
        clone = copy.deepcopy(paragraph._element)
        anchor.addnext(clone)
        anchor = clone
        letters._set_text(_paragraph(clone, document), line)
    return True


def _expand_blocks(document, blocks: dict, context: dict) -> None:
    """{{#name}} … {{/name}} — повторить содержимое для каждого элемента."""
    containers = [document.element.body, *document.element.body.iter(qn("w:tc"))]
    for parent in containers:
        while True:
            children = list(parent)
            start = next((i for i, el in enumerate(children)
                          if el.tag == qn("w:p") and
                          (m := PLACEHOLDER.fullmatch(_paragraph(el, document).text.strip()))
                          and m.group(1) == "#"), None)
            if start is None:
                break
            name = PLACEHOLDER.fullmatch(
                _paragraph(children[start], document).text.strip()).group(2)
            end = next((i for i in range(start + 1, len(children))
                        if children[i].tag == qn("w:p") and
                        _paragraph(children[i], document).text.strip()
                        in ("{{/" + name + "}}", "{{ /" + name + " }}")), None)
            if end is None:
                raise ValueError(f"В шаблоне не закрыт блок {{{{#{name}}}}}")
            protos = children[start + 1:end]
            closing = children[end]
            for item in blocks.get(name) or []:
                scope = {**context, **item}
                for proto in protos:
                    clone = copy.deepcopy(proto)
                    closing.addprevious(clone)
                    if clone.tag == qn("w:p"):
                        if not _fill(_paragraph(clone, document), scope, document):
                            parent.remove(clone)
            for element in [children[start], *protos, closing]:
                parent.remove(element)


def _date_parts(value: str | None) -> tuple[str, str, str]:
    day = date.today()
    if value:
        try:
            day = date.fromisoformat(value)
        except ValueError:
            pass
    return f"{day.day:02d}", letters.MONTHS_GENITIVE[day.month - 1], str(day.year)


def render(act: dict, settings: ActSettings) -> bytes:
    if not settings.ready:
        raise FileNotFoundError("Шаблон акта не найден: " + (settings.template or
                                "в конфиге не указан acts.template"))
    document = Document(settings.template)
    members = [Member.from_dict(m) for m in act.get("members") or []]
    members = [m for m in members if m.org or m.name or m.position]
    day, month, year = _date_parts(act.get("date"))
    context = {
        "organization": act.get("organization") or settings.organization,
        "number": str(act.get("number") or ""),
        "day": day, "month": month, "year": year,
        "body": _body(act.get("body")),
        "appendix": _tidy(act.get("appendix")),
        "basis": _tidy(act.get("basis")),
    }
    blocks = {
        "commission": [{"line": m.commission_line()} for m in members],
        "signatures": [m.signature() for m in members],
    }
    _expand_blocks(document, blocks, context)
    for paragraph in list(_all_paragraphs(document)):
        if not _fill(paragraph, context, document):
            _remove(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _body(text) -> str:
    """Строки текста: пустые выкидываем, первая продолжает «…в том, что»."""
    lines = [" ".join(line.split()) for line in str(text or "").splitlines()]
    return "\n".join(line for line in lines if line)


def file_name(number, subject: str = "") -> str:
    """Как в папке: «Акт № 31 Акт передачи помещений (подвал).doc»."""
    subject = letters.safe_name(subject or "")[:letters.MAX_SUBJECT_IN_NAME]
    name = f"Акт № {number}" + (f" {subject}" if subject else "")
    return letters.safe_name(name) + ".docx"


# — пакет ———————————————————————————————————————————————————————

PACKED = ("body", "subject", "appendix", "basis")


def parties(act: dict) -> list[str]:
    """Стороны по порядку появления: на каждую в форме свой список."""
    found: list[str] = []
    for key in PACKED:
        for name in RE_PARTY.findall(act.get(key) or ""):
            name = name.strip()
            if name and name not in found:
                found.append(name)
    return found


def _put_parties(text: str, chosen: dict) -> str:
    # не выбрали — прочерк, как и в остальных местах, которые вписывают руками
    return RE_PARTY.sub(lambda m: (chosen.get(m.group(1).strip()) or "").strip()
                        or "ООО «___»", text or "")


def fields(act: dict) -> list[str]:
    """Поля пакета по порядку появления — это и есть колонки таблицы."""
    found: list[str] = []
    for key in PACKED:
        for name in RE_FIELD.findall(act.get(key) or ""):
            name = name.strip()
            if name and name not in found:
                found.append(name)
    return found


def parse_rows(text: str) -> list[list[str]]:
    """Строки пакета. Из Excel колонки приходят через табуляцию, руками
    удобнее через точку с запятой."""
    rows = []
    for line in (text or "").splitlines():
        if not line.strip():
            continue
        cells = line.split("\t") if "\t" in line else line.split(";")
        rows.append([cell.strip() for cell in cells])
    return rows


def _put(text: str, values: dict) -> str:
    filled = RE_FIELD.sub(lambda m: values.get(m.group(1).strip(), ""), text or "")
    return re.sub(r"[ \t]{2,}", " ", filled).replace(" .", ".").strip()


def build(act: dict, settings: ActSettings) -> list[tuple[str, bytes]]:
    """Один акт или пакет: по строке на акт, номера подряд."""
    chosen = act.get("parties") or {}
    act = {**act, **{key: _put_parties(act.get(key), chosen)
                     for key in PACKED if act.get(key)}}
    names = fields(act)
    rows = act.get("rows")
    if isinstance(rows, str):
        rows = parse_rows(rows)
    rows = [r for r in rows or [] if any(str(c).strip() for c in r)]

    start = act.get("number")
    try:
        start = int(str(start).strip())
    except (TypeError, ValueError):
        start = None

    if not names:
        number = start if start is not None else act.get("number") or ""
        return [(file_name(number, act.get("subject")),
                 render({**act, "number": number}, settings))]
    if not rows:
        missing = ", ".join("{" + n + "}" for n in names)
        raise ValueError(f"В тексте есть поля {missing}, а строк пакета нет: "
                         "заполните таблицу или уберите фигурные скобки")
    if start is None:
        raise ValueError("Для пакета нужен номер первого акта числом")

    files = []
    for i, row in enumerate(rows):
        values = {name: (row[j] if j < len(row) else "") for j, name in enumerate(names)}
        single = {**act, "number": start + i}
        for key in PACKED:
            single[key] = _put(act.get(key) or "", values)
        files.append((file_name(start + i, single.get("subject")),
                      render(single, settings)))
    return files
